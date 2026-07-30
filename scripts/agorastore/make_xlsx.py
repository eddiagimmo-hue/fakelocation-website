import json
from datetime import datetime
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.drawing.image import Image as XLImage

with open('final.json', encoding='utf-8') as f:
    data = json.load(f)


def end_dt(d):
    return datetime.fromisoformat(d['end_date'])


for d in data:
    d['_end_dt'] = end_dt(d)
    d['_year'] = d['_end_dt'].year

data.sort(key=lambda d: (-(d.get('population') or 0), d['_end_dt']))

wb = Workbook()
wb.remove(wb.active)

headers = [
    "Miniature",
    "Date de fin de vente",
    "Type de local",
    "Code postal",
    "Ville",
    "Habitants",
    "Prix de la mise en vente (EUR)",
    "Prix au m2 (EUR/m2)",
    "Surface (m2)",
    "Surface parcelle exterieure (m2)",
    "URL de l'annonce",
]

header_font = Font(name="Arial", bold=True, color="FFFFFF")
header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
normal_font = Font(name="Arial")

widths = [18, 18, 22, 14, 24, 14, 24, 18, 16, 22, 70]
THUMB_ROW_HEIGHT = 62  # points, ~ 83 px


def write_sheet(ws, rows):
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for row_idx, d in enumerate(rows, start=2):
        type_local = " / ".join(d['types'])

        thumb_path = d.get('thumbnail_path')
        if thumb_path:
            img = XLImage(thumb_path)
            img.width = d.get('thumbnail_width', 120)
            img.height = d.get('thumbnail_height', 80)
            ws.add_image(img, f"A{row_idx}")
        ws.row_dimensions[row_idx].height = THUMB_ROW_HEIGHT

        date_cell = ws.cell(row=row_idx, column=2, value=d['_end_dt'].replace(tzinfo=None))
        date_cell.font = normal_font
        date_cell.number_format = 'DD/MM/YYYY'
        ws.cell(row=row_idx, column=3, value=type_local).font = normal_font
        ws.cell(row=row_idx, column=4, value=d.get('postal_code')).font = normal_font
        ws.cell(row=row_idx, column=5, value=d.get('city')).font = normal_font
        ws.cell(row=row_idx, column=6, value=d.get('population')).font = normal_font
        price_cell = ws.cell(row=row_idx, column=7, value=d.get('initial_price'))
        price_cell.font = normal_font
        price_cell.number_format = '#,##0 "EUR"'
        price_m2 = None
        if d.get('initial_price') is not None and d.get('surface_m2'):
            price_m2 = d['initial_price'] / d['surface_m2']
        price_m2_cell = ws.cell(row=row_idx, column=8, value=price_m2)
        price_m2_cell.font = normal_font
        price_m2_cell.number_format = '#,##0 "EUR"'
        surface_cell = ws.cell(row=row_idx, column=9, value=d.get('surface_m2'))
        surface_cell.font = normal_font
        surface_cell.number_format = '#,##0.00'
        parcelle_cell = ws.cell(row=row_idx, column=10, value=d.get('surface_parcelle_m2'))
        parcelle_cell.font = normal_font
        parcelle_cell.number_format = '#,##0.00'
        url_cell = ws.cell(row=row_idx, column=11, value=d['url'])
        url_cell.font = Font(name="Arial", color="0563C1", underline="single")
        url_cell.hyperlink = d['url']

    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:K{len(rows)+1}"


years = sorted(set(d['_year'] for d in data))

# Summary sheet (created first so it appears as the leftmost tab)
ws2 = wb.create_sheet("Resume")
ws2['A1'] = "Ventes immobilieres terminees sans aucune enchere"
ws2['A1'].font = Font(name="Arial", bold=True, size=14)
ws2['A3'] = "Periode analysee"
ws2['B3'] = "01/01/2023 au " + datetime.now().strftime('%d/%m/%Y')
ws2['A4'] = "Site"
ws2['B4'] = "https://www.agorastore-immo.fr"
ws2['A5'] = "Nombre total d'annonces sans enchere"
ws2['B5'] = len(data)
ws2['A6'] = "Tri"
ws2['B6'] = "Par population de la ville (decroissant), puis par date de fin de vente"

ws2['A8'] = "Nombre d'annonces par annee"
ws2['A8'].font = Font(name="Arial", bold=True)
year_counter = Counter(d['_year'] for d in data)
row = 9
for year in years:
    ws2.cell(row=row, column=1, value=str(year)).font = normal_font
    ws2.cell(row=row, column=2, value=year_counter[year]).font = normal_font
    row += 1

row += 1
ws2.cell(row=row, column=1, value="Repartition par type de local").font = Font(name="Arial", bold=True)
row += 1
counter = Counter()
for d in data:
    for t in d['types']:
        counter[t] += 1
for t, c in sorted(counter.items(), key=lambda x: -x[1]):
    ws2.cell(row=row, column=1, value=t).font = normal_font
    ws2.cell(row=row, column=2, value=c).font = normal_font
    row += 1

for col in ['A', 'B']:
    ws2.column_dimensions[col].width = 45

for r in ws2.iter_rows():
    for c in r:
        if c.font is None or c.font.name != "Arial":
            c.font = normal_font

for year in years:
    rows = [d for d in data if d['_year'] == year]
    ws = wb.create_sheet(str(year))
    write_sheet(ws, rows)

wb.save("annonces_sans_enchere_agorastore.xlsx")
print("saved", len(data), "rows across years:", years)
